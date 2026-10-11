"""
Unit Tests for Report Generator Module.
Validates Excel (.xlsx) workbook generation, crossed vehicles tally,
vehicle inventory construction, and Markdown report output.
"""

import io
import pytest
import openpyxl
from modules.report_generator import ReportGenerator, TrafficReportData


def test_extract_crossed_summary():
    events = [
        {"track_id": 1, "class_name": "Car", "pcu_weight": 1.0, "approach": "North"},
        {"track_id": 2, "class_name": "Bus", "pcu_weight": 3.0, "approach": "South"},
        {"track_id": 3, "class_name": "Motorcycle", "pcu_weight": 0.5, "approach": "North"},
        {"track_id": 4, "class_name": "Truck", "pcu_weight": 3.0, "approach": "East"},
    ]
    summary = ReportGenerator.extract_crossed_summary(events)
    assert summary["total_crossed"] == 4
    assert summary["by_class"]["Car"] == 1
    assert summary["by_class"]["Bus"] == 1
    assert summary["by_class"]["Motorcycle"] == 1
    assert summary["by_class"]["Truck"] == 1
    assert summary["by_lane"]["North"] == 2
    assert summary["by_lane"]["South"] == 1
    assert summary["by_lane"]["East"] == 1
    assert summary["total_pcu_crossed"] == 7.5


def test_build_crossed_vehicles_dataframe():
    events = [
        {
            "track_id": 10,
            "class_name": "Car",
            "timestamp_s": 2.45,
            "frame_index": 74,
            "approach": "North",
            "pcu_weight": 1.0,
            "confidence_pct": 94.0,
            "tripwire_name": "Stopline A",
        }
    ]
    df = ReportGenerator.build_crossed_vehicles_dataframe(events)
    assert len(df) == 1
    assert df["Track ID"].iloc[0] == "#10"
    assert df["Vehicle Class"].iloc[0] == "Car"
    assert df["Crossing Time (s)"].iloc[0] == 2.45
    assert df["PCU Weight"].iloc[0] == 1.0
    assert df["Verification Status"].iloc[0] == "VERIFIED CROSSED"


def test_build_vehicle_inventory_dataframe():
    detection_log = [
        {
            "Timestamp (s)": 1.0,
            "Track ID": "#1",
            "Class": "Car",
            "Confidence": "92%",
            "Approach": "North",
            "PCU Weight": 1.0,
            "In ROI": "Yes",
            "Crossed Line": "Yes",
        },
        {
            "Timestamp (s)": 2.0,
            "Track ID": "#1",
            "Class": "Car",
            "Confidence": "96%",
            "Approach": "North",
            "PCU Weight": 1.0,
            "In ROI": "Yes",
            "Crossed Line": "Yes",
        },
        {
            "Timestamp (s)": 1.5,
            "Track ID": "#2",
            "Class": "Bus",
            "Confidence": "88%",
            "Approach": "South",
            "PCU Weight": 3.0,
            "In ROI": "No",
            "Crossed Line": "No",
        },
    ]
    crossed_ids = {1}
    df = ReportGenerator.build_vehicle_inventory_dataframe(detection_log, crossed_ids)
    assert len(df) == 2

    car_row = df[df["Track ID"] == "#1"].iloc[0]
    assert car_row["Crossed Counting Line?"] == "YES"
    assert car_row["Peak Confidence"] == "96%"
    assert car_row["Total Frames Tracked"] == 2

    bus_row = df[df["Track ID"] == "#2"].iloc[0]
    assert bus_row["Crossed Counting Line?"] == "NO"
    assert bus_row["Vehicle Class"] == "Bus"


def test_generate_excel_traffic_report():
    data = TrafficReportData(
        video_name="test_traffic.mp4",
        total_frames_processed=150,
        video_duration_s=5.0,
        fps=30.0,
        resolution="1280x720",
        source_mode="Manual Video Upload",
        backend_model="YOLOv8 Nano (ONNX)",
        total_unique_vehicles=3,
        total_vehicles_crossed=2,
        crossed_events=[
            {"track_id": 1, "class_name": "Car", "timestamp_s": 1.2, "frame_index": 36, "approach": "North", "pcu_weight": 1.0, "confidence_pct": 95.0},
            {"track_id": 2, "class_name": "Truck", "timestamp_s": 2.8, "frame_index": 84, "approach": "East", "pcu_weight": 3.0, "confidence_pct": 92.0},
        ],
        detection_log=[
            {"Timestamp (s)": 1.2, "Track ID": "#1", "Class": "Car", "Confidence": "95%", "Approach": "North", "PCU Weight": 1.0, "In ROI": "Yes", "Crossed Line": "Yes"},
            {"Timestamp (s)": 2.8, "Track ID": "#2", "Class": "Truck", "Confidence": "92%", "Approach": "East", "PCU Weight": 3.0, "In ROI": "Yes", "Crossed Line": "Yes"},
            {"Timestamp (s)": 3.5, "Track ID": "#3", "Class": "Motorcycle", "Confidence": "89%", "Approach": "South", "PCU Weight": 0.5, "In ROI": "No", "Crossed Line": "No"},
        ],
        telemetry_history=[
            {"timestamp": 1.2, "frame": 36, "total_visible": 1, "unique_tracks": 1, "tripwire_crossings": 1, "flow_vph": 1200, "pcu_total": 1.0, "car": 1, "bus": 0, "truck": 0, "motorcycle": 0},
            {"timestamp": 2.8, "frame": 84, "total_visible": 2, "unique_tracks": 2, "tripwire_crossings": 2, "flow_vph": 1400, "pcu_total": 4.0, "car": 1, "bus": 0, "truck": 1, "motorcycle": 0},
        ],
        signal_plan_data={
            "cycle_time_seconds": 65,
            "delay_reduction_pct": 28.2,
            "phases": [
                {"phase_id": "Phase 1: North-South", "green_time_seconds": 35, "controlled_lanes": ["North", "South"], "priority_level": "Standard"},
                {"phase_id": "Phase 2: East-West", "green_time_seconds": 30, "controlled_lanes": ["East", "West"], "priority_level": "Standard"},
            ],
        },
    )

    excel_bytes = ReportGenerator.generate_excel_traffic_report(data)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 2000

    # Load back with openpyxl to verify sheet integrity
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Executive_Summary" in sheet_names
    assert "Crossed_Vehicles_Log" in sheet_names
    assert "All_Vehicles_Inventory" in sheet_names
    assert "Telemetry_Time_Series" in sheet_names
    assert "Webster_Signal_Plan" in sheet_names

    ws_crossed = wb["Crossed_Vehicles_Log"]
    # Check that rows exist
    assert ws_crossed.max_row >= 6  # Title + headers + 2 crossed rows + totals


def test_generate_markdown_traffic_report():
    data = TrafficReportData(
        video_name="demo_stream.mp4",
        total_frames_processed=100,
        video_duration_s=3.33,
        total_unique_vehicles=2,
        total_vehicles_crossed=1,
        crossed_events=[
            {"track_id": 1, "class_name": "Car", "pcu_weight": 1.0, "approach": "North"},
        ],
    )
    md = ReportGenerator.generate_markdown_traffic_report(data)
    assert "# 🚦 Smart Traffic Management" in md
    assert "Total Vehicles Crossed Line/Cordon" in md
    assert "1 vehicles" in md
